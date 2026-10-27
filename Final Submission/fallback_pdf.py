"""Render our authored PowerPoint shapes and charts to an offline PDF.

Used when LibreOffice is unavailable on a GitHub runner. This is a separate
rendering of the same source PPTX: fonts and some chart details may differ.
"""
import io
from pathlib import Path
import numpy as np
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.text import PP_ALIGN
from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.lib.utils import ImageReader
EMU=914400

def rgb(obj,default):
    try:
        if obj.type is not None and obj.rgb is not None:
            raw=str(obj.rgb)
            return tuple(int(raw[i:i+2],16)/255 for i in (0,2,4))
    except Exception:pass
    return default

def plot_image(c):
    data=[np.array([float(v) if v is not None else 0 for v in s.values],dtype=float) for s in c.series]
    if not data:return None
    names=[s.name for s in c.series]
    try:labels=[str(cat.label) for cat in c.plots[0].categories]
    except Exception:labels=[str(i) for i in range(len(data[0]))]
    clr=["#125449","#C28E52","#BE7460","#8DAE9D"]
    fig,ax=plt.subplots(figsize=(7.7,4.4),dpi=150)
    fig.patch.set_facecolor("white");ax.set_facecolor("white")
    mode=c.chart_type
    if mode==XL_CHART_TYPE.LINE:
        for i,v in enumerate(data):
            ax.plot(range(len(v)),v,color=clr[i%4],linewidth=2.3,label=names[i])
        step=max(1,len(labels)//6)
        ax.set_xticks(list(range(0,len(labels),step)))
        ax.set_xticklabels(labels[::step],fontsize=9)
        ax.grid(axis="y",alpha=.14)
    elif mode==XL_CHART_TYPE.DOUGHNUT:
        ax.pie(data[0],wedgeprops={"width":.4,"edgecolor":"white"},startangle=90,colors=clr)
        ax.legend(labels,loc="lower center",bbox_to_anchor=(.5,-.1),ncol=3,frameon=False,fontsize=9)
    elif mode==XL_CHART_TYPE.BAR_STACKED:
        offset=np.zeros(len(data[0]))
        for i,v in enumerate(data):
            ax.barh(range(len(v)),v,left=offset,color=clr[i%4],label=names[i])
            offset+=v
        ax.set_yticks(range(len(labels)));ax.set_yticklabels(labels,fontsize=10);ax.invert_yaxis()
        ax.grid(axis="x",alpha=.14)
    elif mode==XL_CHART_TYPE.BAR_CLUSTERED:
        n=len(data);width=.72/n
        for i,v in enumerate(data):
            xx=np.arange(len(v))+(i-(n-1)/2)*width
            ax.barh(xx,v,height=width,color=clr[i%4],label=names[i])
        ax.set_yticks(range(len(labels)));ax.set_yticklabels(labels,fontsize=10)
        ax.invert_yaxis();ax.grid(axis="x",alpha=.14)
    else:
        n=len(data);width=.72/n
        for i,v in enumerate(data):
            xx=np.arange(len(v))+(i-(n-1)/2)*width
            ax.bar(xx,v,width=width,color=clr[i%4],label=names[i])
        ax.set_xticks(range(len(labels)));ax.set_xticklabels(labels,fontsize=10)
        ax.grid(axis="y",alpha=.14)
    ax.spines[["top","right","left"]].set_visible(False)
    if len(data)>1 and mode!=XL_CHART_TYPE.DOUGHNUT:
        ax.legend(loc="lower center",bbox_to_anchor=(.5,-.22),ncol=3,frameon=False,fontsize=10)
    fig.tight_layout()
    mem=io.BytesIO();fig.savefig(mem,format="png",dpi=150,bbox_inches="tight",pad_inches=.06)
    plt.close(fig);mem.seek(0)
    return mem

def split_words(text,font,pt,width):
    words=text.split();result=[];line=""
    for word in words:
        trial=(line+" "+word).strip()
        if line and pdfmetrics.stringWidth(trial,font,pt)>width:
            result.append(line);line=word
        else:line=trial
    if line:result.append(line)
    return result or [""]

def convert(input_pptx,output_pdf):
    ppt=Presentation(input_pptx);W=ppt.slide_width/EMU*72;H=ppt.slide_height/EMU*72
    fonts={"/usr/share/fonts/truetype/lato/Lato-Regular.ttf":"Lato",
           "/usr/share/fonts/truetype/lato/Lato-Bold.ttf":"LatoB"}
    for path,name in fonts.items():
        if Path(path).exists():pdfmetrics.registerFont(TTFont(name,path))
    pdf=canvas.Canvas(str(output_pdf),pagesize=(W,H),pageCompression=1)
    for slide in ppt.slides:
        for sh in slide.shapes:
            x=sh.left/EMU*72;y=H-(sh.top+sh.height)/EMU*72
            w=sh.width/EMU*72;h=sh.height/EMU*72
            if getattr(sh,"has_chart",False):
                im=plot_image(sh.chart)
                if im:pdf.drawImage(ImageReader(im),x,y,width=w,height=h,mask="auto")
                continue
            if sh.shape_type==MSO_SHAPE_TYPE.PICTURE:
                pdf.drawImage(ImageReader(io.BytesIO(sh.image.blob)),x,y,width=w,height=h,mask="auto")
                continue
            if sh.shape_type in (MSO_SHAPE_TYPE.AUTO_SHAPE,MSO_SHAPE_TYPE.FREEFORM):
                try:c=rgb(sh.fill.fore_color,(1,1,1))
                except Exception:c=(1,1,1)
                pdf.setFillColorRGB(*c)
                try:
                    if "ROUNDED" in str(sh.auto_shape_type):
                        pdf.roundRect(x,y,w,h,min(h*.12,6),fill=1,stroke=0)
                    else:pdf.rect(x,y,w,h,fill=1,stroke=0)
                except Exception:pdf.rect(x,y,w,h,fill=1,stroke=0)
            if sh.shape_type==MSO_SHAPE_TYPE.LINE:
                try:c=rgb(sh.line.color,(.83,.85,.83))
                except Exception:c=(.83,.85,.83)
                pdf.setStrokeColorRGB(*c)
                pdf.setLineWidth(1)
                pdf.line(x,y+h,x+w,y)
            if not sh.has_text_frame:continue
            lines=[]
            for paragraph in sh.text_frame.paragraphs:
                if not paragraph.text.strip():continue
                pt=paragraph.font.size.pt if paragraph.font.size else 20
                bold=paragraph.font.bold
                font="LatoB" if bold else "Lato"
                hue=rgb(paragraph.font.color,(.16,.22,.18))
                if font not in pdfmetrics.getRegisteredFontNames():
                    font="Helvetica-Bold" if bold else "Helvetica"
                for part in paragraph.text.split("\n"):
                    lines.extend((line,pt,font,hue,paragraph.alignment)
                        for line in split_words(part,font,pt,w-4))
            if not lines:continue
            total=sum(pt*1.14 for _,pt,*_ in lines)
            baseline=y+h/2+total/2-lines[0][1]
            for line,pt,font,hue,align in lines:
                pdf.setFillColorRGB(*hue);pdf.setFont(font,pt)
                if align==PP_ALIGN.CENTER:xx=x+(w-pdfmetrics.stringWidth(line,font,pt))/2
                elif align==PP_ALIGN.RIGHT:xx=x+w-pdfmetrics.stringWidth(line,font,pt)-2
                else:xx=x+2
                pdf.drawString(xx,baseline,line)
                baseline-=pt*1.14
        pdf.showPage()
    pdf.save()
    print("PASS: source-backed PDF exported from PPTX geometry:",output_pdf)
