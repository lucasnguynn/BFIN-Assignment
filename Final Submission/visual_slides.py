#!/usr/bin/env python3
"""Kappa 17-slide visual presentation built offline from committed inputs.

Cream, forest green, gold and rust follow the team's supplied visual style.
All charts are native and editable; the matching PDF is exported from the PPTX.
Sources are in Notes, not in visible slide footers.
"""
import csv,json,math,subprocess
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches,Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE,MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN,MSO_ANCHOR
from pptx.enum.chart import XL_CHART_TYPE,XL_LEGEND_POSITION
from pptx.chart.data import CategoryChartData

D="125449"; G="C28E52"; S="8DAE9D"; R="BE7460"; W="FFFFFF"
C="F8F5EF"; I="2D4038"; M="65786F"; P="EAF0E9"; L="D4DCD3"
COL=[D,G,R,S]
def color(h):return RGBColor.from_string(h)
def readcsv(path):
 with open(path,encoding="utf-8-sig",newline="") as f:return list(csv.DictReader(f))
def shape(sl,x,y,w,h,fill=W,rounded=False):
 sh=sl.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h))
 sh.fill.solid();sh.fill.fore_color.rgb=color(fill);sh.line.fill.background()
 if rounded:
  try:sh.adjustments[0]=.07
  except Exception:pass
 return sh
def rule(sl,x1,y1,x2,y2,hue=L,width=1.2):
 p=sl.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(x1),Inches(y1),Inches(x2),Inches(y2))
 p.line.color.rgb=color(hue);p.line.width=Pt(width)
def t(sl,content,x,y,w,h,size=21,bold=False,hue=I,align=PP_ALIGN.LEFT):
 tb=sl.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h))
 tf=tb.text_frame;tf.clear();tf.word_wrap=True;tf.vertical_anchor=MSO_ANCHOR.MIDDLE
 tf.margin_left=Inches(.04);tf.margin_right=Inches(.04)
 tf.margin_top=Inches(0);tf.margin_bottom=Inches(0)
 for n,st in enumerate(str(content).split("\n")):
  p=tf.paragraphs[0] if n==0 else tf.add_paragraph()
  p.text=st;p.alignment=align;p.font.name="Lato";p.font.size=Pt(size);p.font.bold=bold;p.font.color.rgb=color(hue)
 return tb
def note(sl,s):
 try:sl.notes_slide.notes_text_frame.text="Source and model notes: "+s
 except Exception:pass
def base(ppt,no,title,subtitle,topic):
 sl=ppt.slides.add_slide(ppt.slide_layouts[6]);shape(sl,0,0,20,11.25,C)
 t(sl,"KAPPA FOUNDATION  /  GROUP 4",.73,.28,10,.38,15,True,D)
 t(sl,topic,13.4,.28,5.7,.38,15,True,D,PP_ALIGN.RIGHT)
 rule(sl,.76,.91,19.23,.91,D,1.5)
 t(sl,title,.78,1.19,18.2,.94,43,True,D)
 t(sl,subtitle,.81,2.16,18.0,.61,21,False,M)
 t(sl,str(no).zfill(2),18.63,10.60,.55,.31,14,True,M,PP_ALIGN.RIGHT)
 return sl
def panel(sl,x,y,w,h,tint=W):return shape(sl,x,y,w,h,tint,True)
def kpi(sl,v,lab,x,y,w,hue=D):
 panel(sl,x,y,w,2.00);t(sl,v,x+.23,y+.22,w-.46,.89,39,True,hue)
 t(sl,lab,x+.25,y+1.19,w-.5,.55,17)
def cchart(sl,kind,categories,labels,values,x,y,w,h,minv=None,maxv=None,legend=True,cols=COL):
 data=CategoryChartData();data.categories=categories
 for lab,vec in zip(labels,values):data.add_series(lab,vec)
 chart=sl.shapes.add_chart(kind,Inches(x),Inches(y),Inches(w),Inches(h),data).chart
 chart.has_legend=legend
 if legend:
  chart.legend.position=XL_LEGEND_POSITION.BOTTOM;chart.legend.include_in_layout=False
  chart.legend.font.name="Lato";chart.legend.font.size=Pt(15)
 try:
  ax=chart.value_axis;ax.has_major_gridlines=True
  if minv is not None:ax.minimum_scale=minv
  if maxv is not None:ax.maximum_scale=maxv
  ax.tick_labels.font.name="Lato";ax.tick_labels.font.size=Pt(15)
  chart.category_axis.tick_labels.font.name="Lato"
  chart.category_axis.tick_labels.font.size=Pt(12)
 except Exception:pass
 try:chart.plots[0].gap_width=65
 except Exception:pass
 for i,ser in enumerate(chart.series):
  hue=cols[i%len(cols)]
  try:ser.format.fill.solid();ser.format.fill.fore_color.rgb=color(hue)
  except Exception:pass
  try:ser.format.line.color.rgb=color(hue);ser.format.line.width=Pt(2.6)
  except Exception:pass
 return chart
def boxtext(sl,heading,description,x,y,w,h):
 panel(sl,x,y,w,h,W)
 t(sl,heading,x+.24,y+.17,w-.48,.48,22,True,D)
 t(sl,description,x+.26,y+.73,w-.52,h-.84,18)

def build_visual(root,out):
 root=Path(root);out=Path(out)
 b={r["Metric"]:r for r in readcsv(root/"bond_funding/outputs/revised/revised_funding_summary.csv")}
 ledger=readcsv(root/"bond_funding/outputs/revised/principal_matching_ledger.csv")
 weights=sorted([(r["security"],float(r["weight_of_growth_sleeve"])) for r in readcsv(root/"portfolio_construction/evidence/decision_weights_confirmed.csv") if float(r["weight_of_growth_sleeve"])>0.00001],key=lambda v:v[1])
 monthly=readcsv(root/"portfolio_construction/Portfolio Final/Sources/Observed Monthly.csv")
 observed=json.loads((root/"portfolio_construction/Portfolio Final/Sources/Observed Summary.json").read_text())
 models={r["policy"]:r for r in readcsv(root/"portfolio_construction/outputs/portfolio_comparison.csv")}
 revised={(r["scenario"],r["policy"]):r for r in readcsv(root/"portfolio_construction/revised_mandate/outputs/portfolio_comparison.csv")}
 uf=float(revised[("Revised M1 Principal matching","Frozen Sample MV")]["equity_utility"])
 ur=float(revised[("Revised M1 Principal matching","Reoptimized Sample MV A3")]["equity_utility"])
 gain=ur-uf
 assert 0<gain<1e-6
 fac=readcsv(root/"factor_analysis/outputs/factor_coefficients.csv")
 capm=next(r for r in fac if r["strategy_or_ticker"]=="GROWTH_FROZEN_SAMPLE_MV" and r["model"]=="CAPM")
 ff3=next(r for r in fac if r["strategy_or_ticker"]=="GROWTH_FROZEN_SAMPLE_MV" and r["model"]=="FF3")
 bonds=float(b["Bond cost"]["Principal matching"])
 equity=float(b["Net equity purchases"]["Principal matching"])
 altbond=float(b["Bond cost"]["Coupon funding"])
 altequity=float(b["Net equity purchases"]["Coupon funding"])
 assert abs(bonds-5784228)<.01 and abs(equity-6209562.43756)<.01
 assert len(weights)==8 and abs(sum(x[1] for x in weights)-1)<1e-8
 assert len(monthly)==36
 ppt=Presentation();ppt.slide_width=Inches(20);ppt.slide_height=Inches(11.25)

 # 01 Opening: initial capital split, no stock imagery or decoration.
 s=base(ppt,1,"","","INVESTMENT RECOMMENDATION")
 shape(s,0,2.10,12.30,6.67,D)
 panel(s,12.66,2.10,6.53,6.67,P)
 t(s,"Where the $12m starts",13.05,2.55,5.75,.60,25,True,D)
 cchart(s,XL_CHART_TYPE.BAR_STACKED,["2022 allocation"],["Bonds","Net equities"],
        [[bonds/1e6],[equity/1e6]],13.07,3.37,5.67,2.66,0,12.2,True,[D,G])
 t(s,f"${bonds/1e6:.3f}m bonds",13.10,6.25,5.77,.50,21,True,D)
 t(s,f"${equity/1e6:.3f}m equities",13.10,6.89,5.76,.50,21,True,G)
 t(s,f"${(12e6-bonds-equity):,.0f} entry fee",13.11,7.56,5.72,.50,18,False,M)
 t(s,"KAPPA\nFOUNDATION",1.03,2.70,11.7,3.05,73,True,W)
 t(s,"Investment recommendation\nand revised mandate",1.08,7.00,11.43,1.17,29,False,C)
 t(s,"The $12 million allocation that puts our payment dates first.",1.05,9.62,17.75,.58,24,True,D)
 note(s,"Case-date $12m reconciliation: principal-matching bond cost $5,784,228; net equity $6,209,562.44; 10bp equity entry fee $6,209.56. All quantities are from archived case calculations; no photographs.")

 # 02 Decision.
 s=base(ppt,2,"Our recommendation","Meet the obligations without changing the frozen stock weights.","EXECUTIVE DECISION")
 panel(s,.80,3.15,11.4,6.91)
 panel(s,12.6,3.14,6.58,6.91,P)
 t(s,"Revised commitments",12.99,3.56,5.85,.62,24,True,D)
 cchart(s,XL_CHART_TYPE.COLUMN_CLUSTERED,["2024","2025","2027"],["USD million"],
        [[2.4,1.2,2.4]],12.96,4.40,5.90,4.29,0,3,False,[D])
 t(s,"$6.0m across three dates",13.03,8.94,5.75,.67,20,True,I)
 shape(s,1.26,3.67,3.15,.54,D,True);t(s,"SELECTED POLICY",1.36,3.71,2.96,.41,15,True,W,PP_ALIGN.CENTER)
 t(s,"Principal matching",1.21,4.42,10.37,.90,47,True,D)
 t(s,"2,400 Beta    1,200 Gamma    2,400 Delta",1.29,5.64,10.14,.53,23,True,I)
 rule(s,1.29,6.39,11.65,6.39)
 t(s,"$6.21m",1.30,6.78,4.43,.80,52,True,D)
 t(s,"4.80%",6.38,6.78,4.56,.80,52,True,G)
 t(s,"Net equity investment",1.33,7.67,4.53,.49,19,False,M)
 t(s,"Required equity CAGR to 2030",6.45,7.67,4.65,.69,18,False,M)
 t(s,"A $396k cash reserve remains after the 2025 commitment.",1.34,8.83,10.12,.71,24,True,I)
 note(s,"Group-selected principal bonds 2,400 Beta, 1,200 Gamma, 2,400 Delta; net equity $6,209,562.44 and 4.7951% growth hurdle. Revised payments: $2.4m in 2024, $1.2m in 2025, $2.4m in 2027; the 2025 retained cash is $396k.")

 # 03 Timetable
 s=base(ppt,3,"The $6m total stays the same","Only the timing of the second commitment has changed.","MANDATE")
 panel(s,.79,3.12,18.40,6.96)
 cchart(s,XL_CHART_TYPE.COLUMN_CLUSTERED,["2024","2025","2027"],["Original","Revised"],[[2.4,0,3.6],[2.4,1.2,2.4]],1.40,3.72,16.44,5.52,0,4,True,[S,D])
 t(s,"Revised: $2.4m in 2024, $1.2m in 2025, $2.4m in 2027.",2.16,9.26,15.86,.51,20,True,D,PP_ALIGN.CENTER)
 note(s,"2022 Original $2.4m in 2024 and $3.6m in 2027, vs revised split. USD millions.")

 # 04 Funding tradeoff
 s=base(ppt,4,"We give up growth capital for liquidity","The coupon-funded alternative has more equities but almost no 2025 reserve.","BOND STRATEGY")
 panel(s,.80,3.12,11.38,6.95)
 t(s,"USD millions committed in 2022",1.27,3.57,10.1,.55,25,True,D)
 cchart(s,XL_CHART_TYPE.BAR_STACKED,["Principal","Coupon"],["Bonds","Equities"],[[bonds/1e6,altbond/1e6],[equity/1e6,altequity/1e6]],1.30,4.37,10.53,4.30,0,12,True,[D,G])
 t(s,"Coupon funding adds $354k for equities.",1.36,8.89,10.11,.63,24,True,R)
 panel(s,12.56,3.12,6.65,6.95,P)
 t(s,"Cash after the 2025 payment",12.98,3.59,5.71,.70,24,True,I)
 t(s,"$396k",13.02,4.68,5.69,.93,53,True,D)
 t(s,"Principal matching",13.07,5.66,5.58,.48,20)
 rule(s,13.04,6.49,18.74,6.49)
 t(s,"Nearly zero",13.04,6.82,5.65,.83,41,True,R)
 t(s,"Coupon funding",13.06,7.75,5.56,.53,20)
 t(s,"Both cover the scheduled payments under our cash-flow assumptions.",13.08,8.63,5.51,.98,17,False,I)
 note(s,"Bond costs Principal 5,784,228 and Coupon 5,429,776.18; equity 6,209,562.44 and 6,563,660.16; no double counting.")

 # 05 Cash flows
 s=base(ppt,5,"The 2025 payment has a cash cushion","Selected bonds mature on time and coupons are retained at zero interest.","BOND CASH FLOW")
 panel(s,.78,3.12,12.35,6.96)
 cchart(s,XL_CHART_TYPE.COLUMN_CLUSTERED,[r["Year"] for r in ledger[:5]],["Retained cash, USD thousands"],[[float(r["Closing cash"])/1000 for r in ledger[:5]]],1.17,3.79,11.53,5.71,0,640,False,[D])
 panel(s,13.40,3.12,5.80,6.95,P)
 panel(s,13.60,3.39,5.41,2.56,W)
 t(s,"2025 principal maturity",13.86,3.77,4.88,.53,20,True,I)
 t(s,"$1.20m",13.83,4.42,4.98,.92,41,True,D)
 t(s,"Gamma redemption covers the payment",13.85,5.29,4.89,.48,15,False,M)
 t(s,"2025 reserve",13.91,6.31,4.77,.49,22,True,I)
 t(s,"$396,000",13.88,6.98,4.93,.79,40,True,D)
 rule(s,13.89,8.15,18.74,8.15)
 t(s,"The cash balance is after paying the $1.2m commitment.",13.93,8.39,4.72,.90,17)
 note(s,"Selected principal-matched bond ledger: retained cash $132k end-2023, $264k end-2024, $396k end-2025, $480k end-2026, $564k end-2027. Gamma principal $1.2m meets the 2025 liability.")

 # 06 PV and duration
 s=base(ppt,6,"Bond pricing still responds to yields","Holding bonds to maturity reduces funding risk, not market-price volatility.","BOND MODEL")
 panel(s,.78,3.12,11.43,6.95)
 t(s,"Price response to a parallel +100bp shock",1.28,3.57,9.92,.55,25,True,D)
 cchart(s,XL_CHART_TYPE.BAR_CLUSTERED,["Principal","Coupon"],["Change, USD thousands"],[[-178.26366,-168.60238]],1.27,4.35,10.42,4.33,-200,0,False,[R])
 t(s,"A valuation shock is not a missed bond payment.",1.28,9.07,9.95,.64,21,True,I)
 panel(s,12.57,3.12,6.63,6.95,P)
 t(s,"Case-date present value",13.07,3.58,5.61,.57,22,True,D)
 panel(s,12.99,4.42,5.79,1.47)
 t(s,"P = Σ CFt / (1 + y)^t",13.24,4.72,5.37,.82,27,True,D,PP_ALIGN.CENTER)
 t(s,"Rate approximation",13.07,6.30,5.47,.50,21,True,I)
 panel(s,12.99,7.03,5.78,1.55)
 t(s,"ΔP/P ≈ −Dmod Δy\n+ ½ Convexity (Δy)^2",13.19,7.25,5.43,1.19,22,True,D,PP_ALIGN.CENTER)
 t(s,"Bond marks are modeled, not historical market prices.",13.07,8.89,5.61,.79,16,False,M)
 note(s,"+100bp bond sleeve price change Principal -178,263.66 and Coupon -168,602.38. Modified duration and convexity approximation, cash flow PV.")

 # 07 Capital reconciliation
 s=base(ppt,7,"Where the full $12 million goes","We show the entry fee separately to keep the allocation accurate.","INITIAL CAPITAL")
 panel(s,.78,3.12,11.60,6.95)
 cchart(s,XL_CHART_TYPE.DOUGHNUT,["Bonds","Net equities","Entry fee"],["USD millions"],[[bonds/1e6,equity/1e6,(12e6-bonds-equity)/1e6]],1.44,3.61,9.91,5.95,legend=True,cols=[D,G,R])
 panel(s,12.69,3.12,6.51,6.95,P)
 t(s,"Net equity purchased",13.10,3.63,5.67,.53,22,True,I)
 t(s,"$6.210m",13.07,4.48,5.86,1.00,53,True,D)
 rule(s,13.10,5.80,18.78,5.80)
 t(s,"Entry fee",13.14,6.27,5.49,.54,23,True,I)
 panel(s,13.10,7.07,5.76,1.33)
 t(s,"E = (12m − Bond Cost) / 1.001",13.29,7.26,5.37,.90,20,True,D,PP_ALIGN.CENTER)
 t(s,"Ten basis points: about $6,210.",13.18,8.79,5.60,.61,18,False,M)
 note(s,"Initial bond 5,784,228 + net equity 6,209,562.44 + fee 6,209.56 = $12m exactly.")

 # 08 Individual holdings
 s=base(ppt,8,"Eight stocks, with strict weight limits","The candidate list spans six sectors; selected stocks span five.","STOCK ALLOCATION")
 panel(s,.79,3.12,12.63,6.96)
 cchart(s,XL_CHART_TYPE.BAR_CLUSTERED,[r[0] for r in weights],["Weight %"],[[r[1]*100 for r in weights]],1.12,3.51,11.45,6.00,0,16,False,[D])
 panel(s,13.73,3.12,5.47,6.95,P)
 t(s,"Stock cap",14.18,3.64,4.50,.50,22,True,I)
 t(s,"15%",14.17,4.24,4.51,.87,50,True,D)
 rule(s,14.17,5.42,18.70,5.42)
 t(s,"Technology cap",14.17,5.75,4.53,.56,22,True,I)
 t(s,"30%",14.16,6.37,4.50,.88,50,True,G)
 rule(s,14.18,7.54,18.69,7.54)
 t(s,"No selected ETFs",14.17,7.93,4.52,.61,22,True,D)
 t(s,"Stock concentration still matters.",14.18,8.78,4.53,.76,17,False,M)
 note(s,"Original exact frozen allocation weights, 15% direct-stock cap, 30% sector cap, 30% ETF, 40% combined ETF.")

 # 09 Return vs volatility
 s=base(ppt,9,"More expected return means more estimated risk","Model statistics are in sample, not observed 2023–2025 returns.","PORTFOLIO MODELS")
 panel(s,.79,3.12,12.04,6.96)
 methods=[("Sample MV","Sample_MV_A3"),("Single Index","Single_Index_MV_A3"),("Min. Variance","Minimum_Variance")]
 ret=[100*float(models[key]["annual_expected_arithmetic_return"]) for _,key in methods]
 vol=[100*float(models[key]["annual_volatility"]) for _,key in methods]
 cchart(s,XL_CHART_TYPE.COLUMN_CLUSTERED,[name for name,_ in methods],["Return %","Volatility %"],[ret,vol],1.34,3.75,11.05,5.66,0,17,True,[D,G])
 panel(s,13.13,3.12,6.07,6.95,P)
 for i,(h,body) in enumerate([("Sample MV","Full covariance from 96 months"),("Single Index","Market-factor covariance assumption"),("Minimum Variance","Reduce estimated risk")]):
  boxtext(s,h,body,13.49,3.55+i*2.03,5.33,1.71)
 note(s,"All three models use the common sample covariance for displayed risk: Sample MV 12.0466%/14.9757%; Single Index 11.9055%/14.4624%; Minimum Variance 10.8572%/13.1470%. Single Index is optimized with its own covariance assumption.")

 # 10 Optimization equation
 s=base(ppt,10,"What the optimizer actually solves","The forecasts change, but the constraints stay comparable.","PORTFOLIO MATH")
 panel(s,.78,3.12,11.42,6.96,P)
 t(s,"Risk-adjusted objective",1.25,3.61,10.05,.53,25,True,D)
 panel(s,1.27,4.45,10.43,1.46)
 t(s,"max  μᵀw − (A/2) wᵀΣw",1.46,4.68,10.03,.93,32,True,D,PP_ALIGN.CENTER)
 t(s,"Investment rules",1.28,6.27,9.90,.52,23,True,I)
 for i,q in enumerate(["Σw = 1; w ≥ 0","Each direct stock ≤ 15%","Direct-stock sector ≤ 30%","Each ETF ≤ 30%; combined ETFs ≤ 40%"]):
  t(s,"• "+q,1.50,6.93+i*.62,9.94,.52,19,False,I)
 panel(s,12.58,3.12,6.61,6.96)
 t(s,"Expected-return blend",13.02,3.60,5.68,.55,24,True,D)
 panel(s,12.99,4.46,5.78,1.62,P)
 t(s,"μ = 0.25 μsample\n+ 0.75 μVTI prior",13.25,4.72,5.31,1.12,24,True,D,PP_ALIGN.CENTER)
 rule(s,13.06,6.44,18.70,6.44)
 t(s,"96 months",13.05,6.77,5.60,.84,49,True,G)
 t(s,"Jan 2015 – Dec 2022",13.07,7.75,5.58,.51,20,True,I)
 t(s,"Foundation utility also scales by the equity share.",13.09,8.61,5.55,.83,17,False,M)
 note(s,"Representative mean/variance objective; foundation utility uses E/12m; Sample MV risk aversion A=3, A=6 tested.")

 # 11 Foundation utility
 s=base(ppt,11,"Re-optimization delivers a tiny modeled gain","Keeping the frozen shares avoids overreacting to unstable forecasts.","RISK PREFERENCE")
 panel(s,.78,3.12,11.42,6.95)
 t(s,"Foundation equity contribution to utility",1.23,3.60,10.20,.53,25,True,D)
 panel(s,1.25,4.39,10.46,1.42,P)
 t(s,"U = F μp − (A/2) F² σp²",1.52,4.64,9.90,.94,30,True,D,PP_ALIGN.CENTER)
 t(s,"A = 3, revised Principal Matching",1.26,6.27,10.03,.48,21,True,I)
 for j,(nm,u,hue) in enumerate([("Frozen weights",uf,D),("Reoptimized",ur,G)]):
  y=6.93+j*.86
  t(s,nm,1.29,y-.04,4.43,.52,20,True,I)
  t(s,f"{u:.9f}",6.18,y-.05,5.09,.56,23,True,hue,PP_ALIGN.RIGHT)
 rule(s,1.36,8.60,11.30,8.60)
 t(s,f"Gain: +{gain:.11f} in utility",1.45,8.85,10.02,.62,22,True,D,PP_ALIGN.CENTER)
 panel(s,12.56,3.12,6.63,6.95,P)
 t(s,"Reconsider at higher risk aversion",12.99,3.62,5.63,.68,23,True,D)
 t(s,"A* ≈ 8.94",12.96,4.62,5.74,1.04,50,True,G)
 t(s,"Fixed-policy comparison only",13.05,5.79,5.61,.55,18,False,M)
 rule(s,13.01,6.89,18.73,6.89)
 t(s,"Reassess and re-optimize",13.01,7.29,5.52,.68,23,True,D)
 t(s,"Do not automatically buy the old Minimum Variance portfolio.",13.03,8.26,5.56,.95,18,False,I)
 note(s,f"Source-derived utility frozen {uf:.12f}, reoptimized {ur:.12f}, gain {gain:.12g}. Bars cannot communicate this tiny utility difference faithfully. A*=8.938687 is a pairwise fixed-policy comparison.")

 # 12 CAPM and FF3
 s=base(ppt,12,"FF3 fits the historical data better","A positive estimated intercept is not proof of future stock-picking skill.","FACTOR EVIDENCE")
 panel(s,.79,3.12,9.61,6.95)
 t(s,"Adjusted R squared",1.22,3.59,8.36,.53,25,True,D)
 cchart(s,XL_CHART_TYPE.BAR_CLUSTERED,["CAPM","FF3"],["Adjusted R²"],[[float(capm["r2_adj"]),float(ff3["r2_adj"])]],1.25,4.32,8.51,3.67,0,1,False,[D])
 t(s,"Excess Rp = α + βm excess Rm + ε",1.28,8.23,8.77,.50,20,False,I)
 t(s,"FF3 adds SMB and HML.",1.28,8.89,8.73,.48,21,True,D)
 panel(s,10.73,3.12,8.47,6.95,P)
 t(s,"FF3 factor loadings",11.16,3.58,7.41,.55,25,True,D)
 cchart(s,XL_CHART_TYPE.BAR_CLUSTERED,["Market β","SMB","HML"],["Exposure"],[[float(ff3["Mkt_RF"]),float(ff3["SMB"]),float(ff3["HML"])]],11.13,4.34,7.42,3.77,-.6,1,False,[G])
 rule(s,11.19,8.45,18.68,8.45)
 t(s,"Monthly FF3 α ≈ +0.569% (in sample)",11.17,8.73,7.63,.70,21,True,D)
 note(s,"Factor coefficients from n=96 OLS with HAC sensitivity, CAPM r2adj .784986, FF3 .843652, market .866715, SMB -.42649, HML .0101.")

 # 13 Risk and 2030 hurdle
 s=base(ppt,13,"Funding commitments does not eliminate market losses","The 2030 goal also depends on inflation and equity growth.","STRESS & GOAL")
 panel(s,.80,3.12,11.13,6.95)
 t(s,"Combined shock: stocks −25%, yields +100bp",1.30,3.57,10.02,.53,24,True,D)
 cchart(s,XL_CHART_TYPE.BAR_CLUSTERED,["Principal","Coupon"],["Fund loss %"],[[14.42,15.08]],1.32,4.28,10.18,4.24,0,17,False,[R])
 t(s,"A one-time shock is not a rolling 12-month return.",1.32,9.04,10.02,.61,20,True,I)
 panel(s,12.29,3.12,6.90,6.95,P)
 t(s,"2030 nominal target",12.76,3.58,5.60,.54,22,True,D)
 t(s,"$9.596m",12.73,4.46,5.83,.97,50,True,D)
 t(s,"Assuming 2.3% inflation",12.75,5.52,5.52,.50,18,False,M)
 rule(s,12.72,6.28,18.76,6.28)
 t(s,"Equity CAGR needed",12.75,6.59,5.59,.53,21,True,I)
 t(s,"4.80%",12.73,7.25,5.83,.98,52,True,G)
 panel(s,12.71,8.61,5.95,.95)
 t(s,"r* = [(Target − Cash) / E]^(1/8) − 1",12.86,8.74,5.57,.64,16,True,D,PP_ALIGN.CENTER)
 note(s,"Target 8m*1.023^8=9,596,106.44; principal retained cash 564k; stress fund impact 14.42%; coupon 15.08%.")

 # 14 Actual growth in equity sleeve
 s=base(ppt,14,"The selected shares grew, but VTI did better","Original $6.229m equity base, not revised $6.210m. Monthly returns, 2023–25.","FROZEN EVALUATION")
 panel(s,.79,3.12,18.41,6.95)
 months=[r["month"] for r in monthly]
 stocks=[float(r["growth_value"])/1e6 for r in monthly]
 vti=[float(r["VTI_value"])/1e6 for r in monthly]
 visible_months=[m if i%6==0 or i==len(months)-1 else "" for i,m in enumerate(months)]
 cchart(s,XL_CHART_TYPE.LINE,visible_months,["Selected shares","VTI"],[stocks,vti],1.19,3.54,17.63,5.91,5,12.5,True,[D,G])
 note(s,"Original frozen policy equity base $6,229,246.75, not revised principal $6,209,562.44. Actual archived Yahoo stock prices, 752 common dates, 36 month-end observations, 10bp initial fee, unchanged shares. Not observed whole-fund returns.")

 # 15 Evaluation
 s=base(ppt,15,"Selected shares underperformed the VTI benchmark","Positive absolute returns do not prove an equity-selection advantage.","RESULTS & LIMITATIONS")
 panel(s,.80,3.12,11.39,6.96)
 gg=float(observed["policy"]["CAGR_after_initial_cost"])*100
 vv=float(observed["VTI"]["CAGR_after_initial_cost"])*100
 t(s,"CAGR after entry fee",1.29,3.61,9.71,.55,25,True,D)
 cchart(s,XL_CHART_TYPE.BAR_CLUSTERED,["Selected","VTI"],["CAGR %"],[[gg,vv]],1.34,4.36,10.36,4.19,0,26,False,[D])
 t(s,f"VTI outperformed by {vv-gg:.2f} percentage points",1.30,9.05,10.28,.59,23,True,R)
 panel(s,12.55,3.12,6.65,6.96,P)
 t(s,"Observed risk",13.01,3.60,5.48,.49,25,True,D)
 t(s,"11.80% vs 12.60%",13.02,4.58,5.54,.67,25,True,I)
 t(s,"Annualized volatility",13.04,5.29,5.43,.49,18,False,M)
 rule(s,13.01,6.27,18.73,6.27)
 t(s,"6.91%",13.01,6.63,5.55,.85,50,True,G)
 t(s,"Annualized tracking error",13.05,7.54,5.50,.50,18,False,I)
 t(s,"The future test did not change our 2022 decision.",13.05,8.48,5.51,.90,18,True,D)
 note(s,"Observed growth sleeve CAGR Selected 17.6864%, VTI 22.4469%; volatility 11.80/12.60; tracking error 6.9062.")

 # 16 Governance
 s=base(ppt,16,"Implementation requires monitoring and discipline","Our proposed review rules still require trustee agreement.","GOVERNANCE")
 panel(s,.80,3.12,18.38,6.95)
 for i,(yr,amt) in enumerate([("2024","$2.4m"),("2025","$1.2m"),("2027","$2.4m")]):
  x=1.42+i*5.72;panel(s,x,3.64,4.67,1.38,P)
  t(s,yr,x+.24,3.86,1.51,.53,24,True,D)
  t(s,amt,x+2.08,3.84,2.32,.58,30,True,G,PP_ALIGN.RIGHT)
  if i<2:rule(s,x+4.68,4.33,x+5.62,4.33,D,2)
 t(s,"Rolling withdrawal-adjusted review levels",1.43,5.53,16.8,.54,25,True,D)
 panel(s,1.43,6.46,7.63,2.30,P)
 t(s,"12% trailing loss",1.78,6.81,6.80,.63,31,True,G)
 t(s,"Warning: review cash and assumptions",1.79,7.65,6.77,.54,18)
 panel(s,10.07,6.46,7.61,2.30,P)
 t(s,"15% trailing loss",10.43,6.81,6.73,.63,31,True,R)
 t(s,"Stronger trustee review",10.43,7.65,6.73,.54,18)
 t(s,"Rolling TWR = Π(1 + withdrawal-adjusted monthly return) − 1",1.45,9.28,16.58,.57,20,True,D)
 note(s,"Group-proposed (not trustee-authorized) 12% early warning and 15% review based on rolling withdrawal-adjusted returns. Revised monitor uses case YTM modeled bonds, not historical traded bond prices.")

 # 17 Board decision
 s=base(ppt,17,"One clear recommendation, with risks disclosed","What we would ask the trustees to approve.","BOARD DECISION")
 panel(s,.79,3.13,8.32,6.94,W)
 t(s,"2030 goal under equity scenarios",1.16,3.52,7.61,.55,22,True,D)
 goal=8.0*(1.023**8)
 scenario_rates=[0.0,0.04795107,0.08,0.12]
 cash_at_goal=564000.0
 assert abs(float(ledger[4]["Closing cash"])-cash_at_goal)<0.01
 scenario_total=[(equity*(1+r)**8+cash_at_goal)/1e6 for r in scenario_rates]
 cchart(s,XL_CHART_TYPE.COLUMN_CLUSTERED,["0%","4.80%","8%","12%"],
        ["Fund total, $m","2030 target, $m"],
        [scenario_total,[goal]*4],1.15,4.22,7.54,4.44,0,17,True,[D,G])
 t(s,"Return scenarios, not forecasts",1.22,8.94,7.54,.45,18,True,D)
 t(s,"2.3% inflation | unspent bond cash at 0%",1.22,9.43,7.45,.41,15,False,M)
 panel(s,9.51,3.13,9.68,6.94,P)
 for i,(h,body) in enumerate([("Fund the promises","Principal-match every revised payment."),
    ("Keep the original equity choice","Retain Sample MV weights without hindsight."),
    ("Monitor and reconsider","Use cash needs and rolling losses to review.")]):
  y=3.64+1.86*i
  shape(s,10.03,y,.72,.72,D if i<2 else G,True)
  t(s,str(i+1).zfill(2),10.09,y+.06,.60,.59,20,True,W,PP_ALIGN.CENTER)
  t(s,h,10.97,y-.03,7.48,.54,25,True,D)
  t(s,body,10.98,y+.69,7.51,.65,18,False,I)
  if i<2:rule(s,10.02,y+1.47,18.65,y+1.47)
 t(s,"The 2030 goal is uncertain. Human review and university acceptance remain separate evidence.",10.07,9.35,8.40,.62,16,False,M)
 note(s,"2030 equity scenario sensitivity holds bond cash $564,000 at zero reinvestment: total=net equity (1+r)^8 plus residual cash. 2022-dollar $8m goal inflates at 2.3% for 8 years to $9.596m nominal. These are modeled illustrations, not return forecasts. Funding remains a group proposal; no trustee or LMS approval is implied.")

 assert len(ppt.slides)==17
 count=sum(int(getattr(sh,"has_chart",False)) for sl in ppt.slides for sh in sl.shapes)
 assert count>=15,count
 assert not any(sh.shape_type==13 for slide in ppt.slides for sh in slide.shapes),"No picture shapes permitted"
 ppt.core_properties.title="Kappa Foundation | Group investment presentation"
 ppt.core_properties.author="Kappa Foundation Group 4"
 ppt.save(out/"Presentation.pptx")
 from zipfile import ZipFile
 with ZipFile(out/"Presentation.pptx") as z:
  assert not any(name.startswith("ppt/media/") for name in z.namelist()),"Unexpected raster/image media"
 from fallback_pdf import convert
 convert(out/"Presentation.pptx",out/"Presentation.pdf")
 from pypdf import PdfReader
 assert len(PdfReader(str(out/"Presentation.pdf")).pages)==17
 print(f"PASS: 17 slides, {count} native editable charts, zero photographs, matching PDF")
if __name__=="__main__":
 build_visual(Path(__file__).resolve().parents[1],Path(__file__).resolve().parent)
