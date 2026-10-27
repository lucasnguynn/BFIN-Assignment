#!/usr/bin/env python3
"""Produce a clean, offline and source-checked group submission."""
from pathlib import Path
import csv, json, hashlib, zipfile, re, html
ROOT=Path(__file__).resolve().parents[1]
HERE=Path(__file__).resolve().parent
def csvrows(p):
    with open(p,encoding="utf-8-sig",newline="") as f:return list(csv.DictReader(f))
def plain(t):
    return re.sub(r"\[([^\]]+)\]\([^)]+\)",r"\1",t).replace("*","").replace(chr(96),"")
def blocks():
    out=[]
    for raw in (HERE/"Committee Memo.md").read_text().splitlines():
        s=raw.strip()
        if not s:continue
        if s.startswith("# "):out.append(("title",plain(s[2:])));continue
        if s.startswith("## "):out.append(("section",plain(s[3:])));continue
        if s.startswith("|"):
            cells=[plain(t.strip()) for t in s.strip("|").split("|")]
            if cells and not all(set(c).issubset(set("-: ")) for c in cells):out.append(("table",cells))
            continue
        out.append(("text",plain(s)))
    return out

def verify_sources():
    f={r["Metric"]:r for r in csvrows(ROOT/"bond_funding/outputs/revised/revised_funding_summary.csv")}
    assert abs(float(f["Bond cost"]["Principal matching"])-5784228)<.01
    assert abs(float(f["Net equity purchases"]["Principal matching"])-6209562.437562439)<.01
    assert abs(float(f["Required CAGR"]["Principal matching"])-.04795107162727974)<1e-10
    r=json.loads((ROOT/"portfolio_construction/Portfolio Final/Sources/Observed Summary.json").read_text())
    assert abs(float(r["policy"]["CAGR_after_initial_cost"])-.176864)<.0001
    assert abs(float(r["VTI"]["CAGR_after_initial_cost"])-.224469)<.0001
    qa=json.loads((ROOT/"quality_assurance/qa_report.json").read_text())
    assert qa["counts"]["FAIL"]==0

def memo():
    from docx import Document
    from docx.shared import Inches,Pt
    from reportlab.lib.pagesizes import A4
    from reportlab.lib import colors
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.platypus import SimpleDocTemplate,Paragraph,PageBreak,Table,TableStyle,Spacer
    from pypdf import PdfReader
    doc=Document();sec=doc.sections[0]
    sec.left_margin=sec.right_margin=Inches(.65)
    sec.top_margin=sec.bottom_margin=Inches(.62)
    style=doc.styles["Normal"];style.font.name="Aptos";style.font.size=Pt(9)
    style.paragraph_format.space_after=Pt(6)
    styles={
      "title":ParagraphStyle("t",fontName="Helvetica-Bold",fontSize=14,leading=18,spaceAfter=9),
      "section":ParagraphStyle("s",fontName="Helvetica-Bold",fontSize=11,leading=14,spaceAfter=8),
      "text":ParagraphStyle("b",fontName="Helvetica",fontSize=8.55,leading=12.15,spaceAfter=7),
      "cell":ParagraphStyle("c",fontName="Helvetica",fontSize=7.9,leading=10)}
    story=[];sections=0;table=[]
    def flush():
        nonlocal table
        if not table:return
        n=len(table[0])
        cells=[[Paragraph(html.escape(x),styles["cell"]) for x in row] for row in table]
        t=Table(cells,colWidths=[(A4[0]-88)/n]*n,repeatRows=1)
        t.setStyle(TableStyle([
          ("BACKGROUND",(0,0),(-1,0),colors.HexColor("#EEEEEE")),
          ("LINEBELOW",(0,0),(-1,0),.5,colors.gray),
          ("VALIGN",(0,0),(-1,-1),"TOP"),
          ("TOPPADDING",(0,0),(-1,-1),4),
          ("BOTTOMPADDING",(0,0),(-1,-1),4)]))
        story.extend([t,Spacer(1,6)])
        tb=doc.add_table(rows=len(table),cols=n);tb.style="Light Shading Accent 1"
        for i,row in enumerate(table):
            for j,x in enumerate(row):tb.cell(i,j).text=x
        table=[]
    for kind,val in blocks():
        if kind=="table":table.append(val);continue
        flush()
        if kind=="section":
            if sections:story.append(PageBreak());doc.add_page_break()
            sections+=1;doc.add_heading(val,2)
        elif kind=="title":doc.add_heading(val,0)
        else:doc.add_paragraph(val)
        story.append(Paragraph(html.escape(val),styles[kind]))
    flush()
    if sections!=8:raise AssertionError(f"Expected eight memo sections, got {sections}")
    doc.save(HERE/"Committee Memo.docx")
    def footer(c,d):
        c.setFont("Helvetica",7);c.setFillGray(.45)
        c.drawString(44,24,"Kappa Foundation | Committee memo")
        c.drawRightString(A4[0]-44,24,str(d.page))
    SimpleDocTemplate(str(HERE/"Committee Memo.pdf"),pagesize=A4,
        leftMargin=44,rightMargin=44,topMargin=42,bottomMargin=43).build(
        story,onFirstPage=footer,onLaterPages=footer)
    pages=len(PdfReader(str(HERE/"Committee Memo.pdf")).pages)
    if pages>8:raise AssertionError(f"Committee memo has {pages} pages; 8 is the maximum including references")
    return pages

def decisions():
    from reportlab.platypus import SimpleDocTemplate,Paragraph,Spacer
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.pagesizes import A4
    t=ParagraphStyle("entry",fontSize=9.5,leading=14,spaceAfter=12)
    story=[Paragraph("<b>Decision and contributions</b>",t)]
    notes=[
      ("Funding","Group-selected principal matching: 2,400 Beta, 1,200 Gamma, 2,400 Delta."),
      ("Stocks","Original Sample MV weights retained after revised funding."),
      ("Evaluation","The 2023-2025 results are a separate conditional test, not a 2022 model input."),
      ("Human review","The group representative reports unanimous agreement with the recommendation; this does not constitute individually authenticated signatures or independent review of each file."),
      ("Provenance","Original emailed archive-member hashes and university acceptance have not been independently confirmed."),
      ("Academic integrity","Students report choosing investment ideas, methods, assumptions, securities, sources and document design. AI supported code, research, data processing, calculations, QA and editing under human decision control.")
    ]
    for h,b in notes:story.append(Paragraph("<b>"+h+"</b>. "+b,t))
    SimpleDocTemplate(str(HERE/"Decision Record.pdf"),pagesize=A4,
      topMargin=54,leftMargin=48,rightMargin=48).build(story)

def pack():
    analysis=HERE/"Analysis.zip"
    with zipfile.ZipFile(analysis,"w",zipfile.ZIP_DEFLATED,compresslevel=6) as z:
        for p in sorted(ROOT.rglob("*")):
            if not p.is_file():continue
            rel=p.relative_to(ROOT).as_posix()
            if rel.startswith("Final Submission/"):continue
            if any(part in (".git","__pycache__", ".venv",".pytest_cache") for part in p.parts):continue
            if rel=="portfolio_construction/Portfolio Final/Portfolio Files.zip":continue
            if not (rel.startswith(("bond_funding/","portfolio_construction/","factor_analysis/","quality_assurance/")) or
                    rel in ("README.md","requirements.txt","run_pipeline.py")):continue
            z.write(p,rel)
        # These files are necessary to rebuild the final memo and slides, not only the research tables.
        for name in ("Committee Memo.md","build_final.py","visual_slides.py","fallback_pdf.py"):
            z.write(HERE/name,"Final Submission/"+name)
        for asset in sorted((HERE/"assets").glob("*")):
            if asset.is_file():z.write(asset,"Final Submission/assets/"+asset.name)
    with zipfile.ZipFile(analysis) as z:
        assert z.testzip() is None
        assert "Final Submission/build_final.py" in z.namelist()
        assert "Final Submission/visual_slides.py" in z.namelist()
        assert {"Final Submission/assets/Campus.jpg","Final Submission/assets/Study.jpg",
                "Final Submission/assets/Sources.json"}.issubset(set(z.namelist()))
    (HERE/"Read Me.txt").write_text("Kappa Foundation | BFIN WS2026 final submission\n\nRead Committee Memo.pdf and Presentation.pdf first. Editable versions are included.\nAnalysis.zip contains the research, original frozen decision, saved Yahoo and Kenneth French data, Python source, tests, models and generated outputs. It also contains the source scripts for this final memo and slide deck.\nThe revised principal-matching recommendation retains the original Sample MV stock weights. The 2023-2025 observed growth-only comparison uses the original frozen equity budget, not a revised budget.\n\nOFFLINE REBUILD (Python 3.12)\n1. Extract this ZIP, then extract Analysis.zip into the same folder (preserving its directory structure).\n2. In that extracted project root, install packages: python -m pip install -r requirements.txt\n3. Rebuild all workstreams and final documents: python \"Final Submission/build_final.py\" --rebuild-offline\n4. Check the three unit-test suites and QA output under quality_assurance/qa_report.json. The command above runs those tests.\n5. Check regenerated Checksums.txt against the seven deliverables; final ZIP also includes this checksum record.\nNo live download is needed. Retrieval scripts are kept as source evidence but are not called by the offline rebuild.\n\nThe 22 October university email carrying a first draft ZIP was directly checked; it is not proof of LMS or instructor acceptance, and emailed archive-member hashes remain unverified. The group representative reports collective agreement, not verified individual signatures. JNJ/Kenvue issuer treatment is confirmed, while automated Yahoo vendor-price checks are not personal sign-off.\nIllustrative campus and seminar-library photos are included in the deck, with CC0 license, authors, source URLs and SHA-256 in Analysis.zip/Final Submission/assets/Sources.json. They do not depict the fictional Kappa Foundation.\nThe board has not approved the group's proposed 12% warning and 15% review thresholds.\nThe students report directing investment ideas, course methods, securities, assumptions, source choices, feasibility, analysis interpretation and design. AI materially supported coding, discovery/collection of information, data processing, quantitative execution, testing, document production and editing under human oversight. See quality_assurance/Group Review and AI Use.md in Analysis.zip."+"\n",encoding="utf-8")
    items=["Committee Memo.pdf","Committee Memo.docx","Presentation.pdf",
           "Presentation.pptx","Decision Record.pdf","Analysis.zip","Read Me.txt"]
    sha=lambda p:hashlib.sha256(p.read_bytes()).hexdigest()
    (HERE/"Checksums.txt").write_text("".join(sha(HERE/n)+"  "+n+"\n" for n in items),encoding="utf-8")
    final=HERE/"Final Submission.zip"
    with zipfile.ZipFile(final,"w",zipfile.ZIP_DEFLATED,compresslevel=6) as z:
        for name in (*items,"Checksums.txt"):
            z.write(HERE/name,name)
    with zipfile.ZipFile(final) as z:
        assert z.testzip() is None
        assert sorted(z.namelist())==sorted([*items,"Checksums.txt"])
        for name in items:
            assert hashlib.sha256(z.read(name)).hexdigest()==sha(HERE/name)
    return final

def rebuild_offline():
    """Rebuild from saved inputs only, including the reader-facing final documents."""
    import shutil,subprocess,sys,tempfile
    def run(*parts):
        subprocess.run([sys.executable,*map(str,parts)],cwd=ROOT,check=True)
    run("run_pipeline.py")
    run("portfolio_construction/revised_mandate/m1_integration.py",
        "--out","portfolio_construction/revised_mandate/outputs")
    with tempfile.TemporaryDirectory(prefix="kappa-offline-") as tmp:
        staging=Path(tmp)/"observed"
        yahoo=staging/"observed_yahoo_data"
        provisional=staging/"observed_provisional_results"
        yahoo.mkdir(parents=True)
        provisional.mkdir(parents=True)
        source=ROOT/"portfolio_construction/Portfolio Final/Sources"
        for original,target in (
            ("Yahoo Raw.zip",yahoo/"yfinance_evaluation_RAW.zip"),
            ("Observed Prices.csv",yahoo/"daily_adjusted_prices_2023_2025.csv"),
            ("Market Source.json",yahoo/"evaluation_source_record.json"),
            ("Observed Monthly.csv",provisional/"PROVISIONAL observed 2023-2025 monthly growth VTI.csv"),
            ("Observed Summary.json",provisional/"PROVISIONAL observed performance summary.json"),
        ):
            shutil.copyfile(source/original,target)
        run("portfolio_construction/build_portfolio_files.py","--observed",staging)
    # The portfolio acceptance tests require its user-facing ZIP, rebuilt above.
    for suite in ("portfolio_construction/tests","factor_analysis/tests","quality_assurance/tests"):
        run("-m","unittest","discover","-s",suite)
    run("portfolio_construction/post_freeze_evaluation/jnj_vendor_check.py")
    run("portfolio_construction/post_freeze_evaluation/revised_fund_monitor.py",
        "--growth-monthly","portfolio_construction/Portfolio Final/Sources/Observed Monthly.csv",
        "--out","portfolio_construction/outputs/revised_fund_monitor.csv",
        "--policy","principal")

def main():
    verify_sources();pages=memo()
    from visual_slides import build_visual
    build_visual(ROOT,HERE)
    decisions();archive=pack()
    print(json.dumps({"status":"PASS","memo_pages":pages,"slides":17,
        "final_zip_bytes":archive.stat().st_size},indent=2))
if __name__=="__main__":
    import sys
    if len(sys.argv)==2 and sys.argv[1]=="--rebuild-offline":rebuild_offline()
    elif len(sys.argv)>1:raise SystemExit("Usage: build_final.py [--rebuild-offline]")
    main()
