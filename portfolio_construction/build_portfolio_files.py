#!/usr/bin/env python3
"""Make the clean portfolio folder using the saved GitHub evidence."""
import argparse,csv,hashlib,json,re,shutil,zipfile,html
from pathlib import Path
from xml.sax.saxutils import escape
BASE=Path(__file__).resolve().parent
REPO=BASE.parent
FINAL=BASE/"Portfolio Final"
def rows(p):
 with open(p,encoding="utf-8-sig",newline="") as f:return list(csv.DictReader(f))
def sha(p):return hashlib.sha256(Path(p).read_bytes()).hexdigest()
def blocks(text):
 out=[];lines=text.splitlines();i=0
 while i<len(lines):
  s=lines[i].strip();i+=1
  if not s:continue
  if s.startswith("# "):out.append(("title",s[2:]));continue
  if s.startswith("## "):out.append(("heading",s[3:]));continue
  if s.startswith("|"):
   table=[]
   while True:
    vals=[v.strip() for v in s.strip("|").split("|")]
    if not all(re.fullmatch("[-: ]+",v or "") for v in vals):table.append(vals)
    if i>=len(lines) or not lines[i].strip().startswith("|"):break
    s=lines[i].strip();i+=1
   out.append(("table",table));continue
  out.append(("text",s))
 return out
def plain(s):
 return re.sub(r"\*\*([^*]+)\*\*",r"\1",s).replace("&#","")
def documents(out):
 from docx import Document
 from docx.shared import Pt,Inches
 from reportlab.platypus import SimpleDocTemplate,Paragraph,Table,TableStyle,Spacer
 from reportlab.lib.styles import getSampleStyleSheet,ParagraphStyle
 from reportlab.lib.colors import HexColor
 from reportlab.lib.pagesizes import A4
 b=blocks((BASE/"Portfolio Review.md").read_text(encoding="utf8"))
 d=Document();sec=d.sections[0];sec.top_margin=Inches(.60);sec.bottom_margin=Inches(.60)
 sec.left_margin=Inches(.68);sec.right_margin=Inches(.68)
 normal=d.styles["Normal"];normal.font.name="Aptos";normal.font.size=Pt(9)
 normal.paragraph_format.space_after=Pt(4)
 ss=getSampleStyleSheet()
 styles={"title":ParagraphStyle("main",parent=ss["Heading1"],fontName="Helvetica-Bold",fontSize=14,leading=17,spaceAfter=7),
 "heading":ParagraphStyle("section",parent=ss["Heading2"],fontName="Helvetica-Bold",fontSize=10.2,leading=13,spaceBefore=8,spaceAfter=4),
 "text":ParagraphStyle("body2",parent=ss["BodyText"],fontName="Helvetica",fontSize=8.6,leading=12,spaceAfter=6)}
 story=[]
 for typ,v in b:
  if typ=="table":
   t=d.add_table(rows=len(v),cols=len(v[0]));t.style="Light Shading Accent 1"
   for i,row in enumerate(v):
    for j,x in enumerate(row):t.cell(i,j).text=plain(x)
   if len(v)>0:
    usable=(A4[0]-86)/len(v[0])
    z=Table([[Paragraph(html.escape(plain(x)),styles["text"]) for x in row] for row in v],
      colWidths=[usable]*len(v[0]),repeatRows=1)
    z.setStyle(TableStyle([("BACKGROUND",(0,0),(-1,0),HexColor("#eeeeee")),
      ("VALIGN",(0,0),(-1,-1),"TOP"),("BOTTOMPADDING",(0,0),(-1,-1),5),
      ("LINEBELOW",(0,0),(-1,0),.5,HexColor("#999999"))]))
    story.append(z)
  else:
   if typ=="title":d.add_heading(plain(v),0)
   elif typ=="heading":d.add_heading(plain(v),2)
   else:d.add_paragraph(plain(v))
   story.append(Paragraph(html.escape(plain(v)),styles[typ]))
 d.save(out/"Portfolio Report.docx")
 SimpleDocTemplate(str(out/"Portfolio Report.pdf"),pagesize=A4,topMargin=39,bottomMargin=38,
   leftMargin=43,rightMargin=43).build(story)
def xlcol(n):
 s=""
 while n:n,k=divmod(n-1,26);s=chr(65+k)+s
 return s
def workbook(sheets,path):
 """Small valid Office Open XML workbook. No Excel external connections."""
 types=['<?xml version="1.0"?><Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"><Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/><Default Extension="xml" ContentType="application/xml"/><Override PartName="/xl/workbook.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"/><Override PartName="/xl/styles.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.styles+xml"/>']
 types += [f'<Override PartName="/xl/worksheets/sheet{i}.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/>' for i in range(1,len(sheets)+1)]
 wb='<workbook xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"><sheets>'+''.join(f'<sheet name="{escape(name)}" sheetId="{i}" r:id="rId{i}"/>' for i,(name,_) in enumerate(sheets,1))+'</sheets></workbook>'
 rel='<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'+''.join(f'<Relationship Id="rId{i}" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet" Target="worksheets/sheet{i}.xml"/>' for i in range(1,len(sheets)+1))+'<Relationship Id="rIdS" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/styles" Target="styles.xml"/></Relationships>'
 style='<styleSheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><fonts count="2"><font><sz val="10"/><name val="Aptos"/></font><font><b/><sz val="10"/><name val="Aptos"/></font></fonts><fills count="3"><fill><patternFill patternType="none"/></fill><fill><patternFill patternType="gray125"/></fill><fill><patternFill patternType="solid"><fgColor rgb="FFEEEEEE"/></patternFill></fill></fills><borders count="1"><border><left/><right/><top/><bottom/><diagonal/></border></borders><cellStyleXfs count="1"><xf numFmtId="0" fontId="0" fillId="0" borderId="0"/></cellStyleXfs><cellXfs count="2"><xf numFmtId="0" fontId="0" fillId="0" borderId="0" xfId="0"/><xf numFmtId="0" fontId="1" fillId="2" borderId="0" xfId="0"/></cellXfs></styleSheet>'
 with zipfile.ZipFile(path,"w",zipfile.ZIP_DEFLATED) as z:
  z.writestr("[Content_Types].xml","".join(types)+"</Types>")
  z.writestr("_rels/.rels",'<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="xl/workbook.xml"/></Relationships>')
  z.writestr("xl/workbook.xml",wb);z.writestr("xl/_rels/workbook.xml.rels",rel);z.writestr("xl/styles.xml",style)
  for i,(name,vals) in enumerate(sheets,1):
   rr=[]
   for a,row in enumerate(vals,1):
    cc=[]
    for b,x in enumerate(row,1):
     if x is None:continue
     ref=xlcol(b)+str(a);sty=' s="1"' if a==1 else ''
     if isinstance(x,(int,float)):cc.append(f'<c r="{ref}"{sty}><v>{x:.16g}</v></c>')
     else:cc.append(f'<c r="{ref}"{sty} t="inlineStr"><is><t>{escape(str(x))}</t></is></c>')
    rr.append(f'<row r="{a}">{"".join(cc)}</row>')
   z.writestr(f"xl/worksheets/sheet{i}.xml",'<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><sheetViews><sheetView workbookViewId="0"><pane ySplit="1" topLeftCell="A2" activePane="bottomLeft" state="frozen"/></sheetView></sheetViews><cols>'+''.join(f'<col min="{j}" max="{j}" width="{30 if j==1 else 20}" customWidth="1"/>' for j in range(1,max(map(len,vals))+1))+'</cols><sheetData>'+"".join(rr)+'</sheetData></worksheet>')
def create(observed):
 observed=Path(observed);FINAL.mkdir(parents=True,exist_ok=True)
 o=observed/"observed_yahoo_data"
 p=observed/"observed_provisional_results"
 raw=o/"yfinance_evaluation_RAW.zip";date=o/"evaluation_source_record.json"
 prices=o/"daily_adjusted_prices_2023_2025.csv"
 monthly=p/"PROVISIONAL observed 2023-2025 monthly growth VTI.csv"
 summary=p/"PROVISIONAL observed performance summary.json"
 metadata=json.loads(date.read_text());sv=json.loads(summary.read_text())
 if sha(raw)!=metadata["raw_archive_sha256"] or sha(prices)!=metadata["processed_price_csv_sha256"]:raise AssertionError("Yahoo SHA invalid")
 if len(rows(prices))!=752 or len(rows(monthly))!=36:raise AssertionError("Observed data incomplete")
 from sys import path as sys_path
 sys_path.insert(0,str(BASE/"revised_mandate"))
 from m1_integration import run as run_model
 from policy_sensitivity import build as build_goal
 from m1_peer_check import check as check_bonds
 calc=FINAL/"Sources";calc.mkdir(exist_ok=True)
 out=FINAL/".temp_calc";out.mkdir(exist_ok=True)
 run_model(out)
 build_goal(out/"goal.csv")
 check_bonds(out/"cash_check.csv")
 funding=rows(BASE.parent/"bond_funding/outputs/revised/revised_funding_summary.csv")
 allocations=rows(out/"exact_allocations.csv")
 models=rows(out/"portfolio_comparison.csv")
 cash=[]
 for fn in ("principal_matching_ledger.csv","coupon_funding_ledger.csv"):
  for row in rows(BASE.parent/"bond_funding/outputs/revised"/fn):cash.append({"Policy":fn.split("_")[0],**row})
 fac=rows(BASE.parent/"factor_analysis/outputs/factor_coefficients.csv")
 def tab(data):return [list(data[0].keys())]+[[float(v) if v and re.fullmatch(r"-?\d+(\.\d+)?([eE][+-]?\d+)?",v) else v for v in r.values()] for r in data]
 observed_rows=rows(monthly)
 alpha=[r for r in fac if r["strategy_or_ticker"] in ("GROWTH_FROZEN_SAMPLE_MV","MSFT","PG") and r["model"]=="FF3"]
 sources=[
 ["Source","Note"],["2022 decision","Original case decision PDF and original repository code"],
 ["Revised bonds","bond_funding/outputs/revised"],
 ["Yahoo archive","Sources/Yahoo Raw.zip; 752 common trading days, 36 monthly results"],
 ["Observed data","Retrieved after the original 22 October course email; acceptance not independently confirmed"],
 ["Case-date cutoff","2015 to 2022 estimation only; 2023 to 2025 separately evaluated"],
 ["JNJ","Optional tender not elected; human review remains pending"],
 ["Group approval","Final bond-funding choice and peer checks pending"],
 ["AI disclosure","Material assistance in code, source review and editing"]]
 sheets=[
 ("Summary",[["Metric","Selected stocks","VTI"],["After-fee CAGR",float(sv["policy"]["CAGR_after_initial_cost"]),float(sv["VTI"]["CAGR_after_initial_cost"])],["Annual volatility",float(sv["policy"]["volatility_annualized"]),float(sv["VTI"]["volatility_annualized"])],["Maximum monthly drawdown",float(sv["policy"]["maximum_drawdown"]),float(sv["VTI"]["maximum_drawdown"])],["Final equity USD",float(sv["policy"]["terminal_growth_USD"]),float(sv["VTI"]["terminal_growth_USD"])] ]),
 ("Funding",tab(funding)),("Investments",tab(allocations)),("Models",tab(models)),
 ("Cash Flows",tab(cash)),("Risk and Goal",tab(rows(out/"goal.csv"))),
 ("Observed",tab(observed_rows)),("Factors",tab(alpha)),("Sources and Status",sources)]
 workbook(sheets,FINAL/"Results.xlsx");documents(FINAL)
 shutil.copyfile(raw,calc/"Yahoo Raw.zip")
 shutil.copyfile(prices,calc/"Observed Prices.csv")
 shutil.copyfile(date,calc/"Market Source.json")
 shutil.copyfile(monthly,calc/"Observed Monthly.csv")
 shutil.copyfile(summary,calc/"Observed Summary.json")
 # Two simple plots are the only figures kept with the portfolio handoff.
 import matplotlib
 matplotlib.use("Agg")
 import matplotlib.pyplot as plt
 chart_dir=calc/"Charts";chart_dir.mkdir(exist_ok=True)
 dates=[r["month"] for r in observed_rows]
 chosen=[float(r["growth_value"])/1000000 for r in observed_rows]
 benchmark=[float(r["VTI_value"])/1000000 for r in observed_rows]
 differences=[100*float(r["active_monthly"]) for r in observed_rows]
 fig,ax=plt.subplots(figsize=(7.8,3.4))
 ax.plot(range(len(dates)),chosen,label="Selected stocks",linewidth=1.9)
 ax.plot(range(len(dates)),benchmark,label="VTI",linewidth=1.9)
 ax.set_ylabel("Growth sleeve, USD million");ax.set_xlabel("Month")
 ax.set_xticks(range(0,len(dates),6),[dates[i] for i in range(0,len(dates),6)])
 ax.legend(frameon=False);ax.grid(alpha=.15);fig.tight_layout()
 fig.savefig(chart_dir/"Growth vs VTI.png",dpi=160);plt.close(fig)
 fig,ax=plt.subplots(figsize=(7.8,3.0))
 ax.bar(range(len(dates)),differences)
 ax.axhline(0,linewidth=.7,color="black")
 ax.set_ylabel("Monthly active return, %");ax.set_xlabel("Month")
 ax.set_xticks(range(0,len(dates),6),[dates[i] for i in range(0,len(dates),6)])
 ax.grid(axis="y",alpha=.15);fig.tight_layout()
 fig.savefig(chart_dir/"Monthly Difference.png",dpi=160);plt.close(fig)

 shutil.copyfile(BASE/"docs/Kappa Foundation- Initial Decision.pdf",calc/"Original Decision.pdf")
 (calc/"Read Me.txt").write_text("PORTFOLIO SOURCES\n\nRead Portfolio Report.pdf and Results.xlsx first. The Python source, original 2015-2022 returns, revised bond model and French factor data are elsewhere in this GitHub repository. The raw Yahoo responses here are preserved, dated and checksum-verified.\n\nOriginal university email submission is documented, but instructor acceptance and exact emailed ZIP-member hashes are not confirmed. Final funding approval and peer review are reserved for the group.\n\nMaterial AI assistance in coding, checking and editing is disclosed in the report.\n")
 shutil.rmtree(out)
 entries=[(p.relative_to(FINAL).as_posix(),sha(p)) for p in FINAL.rglob("*") if p.is_file() and p.name not in ("Checksums.txt","Portfolio Files.zip")]
 (calc/"Checksums.txt").write_text("".join(f"{h}  {name}\n" for name,h in sorted(entries)))
 with zipfile.ZipFile(FINAL/"Portfolio Files.zip","w",zipfile.ZIP_DEFLATED) as z:
  for path in FINAL.rglob("*"):
   if path.is_file() and path.name!="Portfolio Files.zip":z.write(path,path.relative_to(FINAL))
 assert zipfile.ZipFile(FINAL/"Portfolio Files.zip").testzip() is None
 print("PASS: final portfolio PDF Word Excel, 9 sheets, 36-month results, raw Yahoo SHA and bundle ready")
if __name__=="__main__":
 parser=argparse.ArgumentParser();parser.add_argument("--observed",required=True)
 create(parser.parse_args().observed)
